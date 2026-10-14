# -*- coding: utf-8 -*-
"""
Script: baixar_base_procampo.py
Objetivo: manter BD_PROCAMPO/BASE_PRODUTORES_PROCAMPO_PIRACANJUBA.xlsx igual à versão do SharePoint (a fonte).

Lê o arquivo pelo link de compartilhamento (procampo.base_sharepoint_link no config.yaml) via Microsoft Graph, com as
credenciais TENANT_ID / CLIENT_ID / CLIENT_SECRET de scripts/config/.env. Só baixa quando o SharePoint está mais novo
que a cópia local; a versão anterior vai para BD_PROCAMPO/BACKUPS antes de ser substituída. Qualquer falha deixa a cópia
local intacta (a carga do Procampo segue com ela).

Uso isolado:
  python scripts/functions/baixar_base_procampo.py
"""
from __future__ import annotations

import base64
import os
import shutil
import sys
import tempfile
import zipfile
from datetime import datetime, timezone
from pathlib import Path

import requests
import yaml
from dotenv import load_dotenv

if hasattr(sys.stdout, "reconfigure"):
    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

GRAPH = "https://graph.microsoft.com/v1.0"
TOLERANCIA_S = 60  # diferença de relógio/gravação que não conta como "mais novo"


def detectar_raiz() -> Path:
    return Path(__file__).resolve().parent.parent.parent


def _token(tenant: str, client: str, segredo: str) -> str:
    r = requests.post(
        f"https://login.microsoftonline.com/{tenant}/oauth2/v2.0/token",
        data={"grant_type": "client_credentials", "client_id": client, "client_secret": segredo,
              "scope": "https://graph.microsoft.com/.default"},
        timeout=30,
    )
    if r.status_code != 200:
        raise RuntimeError(f"autenticação no Microsoft Graph falhou (HTTP {r.status_code})")
    return r.json()["access_token"]


def _id_compartilhamento(link: str) -> str:
    """Codificação exigida pela API /shares: 'u!' + base64 url-safe do link, sem '='."""
    return "u!" + base64.urlsafe_b64encode(link.encode("utf-8")).decode("ascii").rstrip("=")


def sincronizar_base_procampo(raiz: Path | None = None) -> bool:
    """Baixa a base do SharePoint se estiver mais nova. Devolve True se atualizou, False se já estava em dia."""
    raiz = raiz or detectar_raiz()
    try:
        from functions.config_projeto import ler_config
    except ImportError:
        from config_projeto import ler_config
    cfg = ler_config(raiz)
    cfg_pro = cfg["procampo"]
    link = cfg_pro.get("base_sharepoint_link")
    if not link:
        raise RuntimeError("procampo.base_sharepoint_link não configurado no config.yaml")

    load_dotenv(raiz / "scripts" / "config" / ".env")
    faltando = [v for v in ("TENANT_ID", "CLIENT_ID", "CLIENT_SECRET") if not os.getenv(v)]
    if faltando:
        raise RuntimeError(f"variáveis ausentes em scripts/config/.env: {', '.join(faltando)}")

    pasta = Path(cfg["caminhos"]["bd_procampo"])
    destino = pasta / cfg_pro["arquivo_base"]

    print("\n" + "=" * 70)
    print("☁️  SINCRONIZANDO A BASE DE PRODUTORES DO PROCAMPO COM O SHAREPOINT")
    print("=" * 70)

    cabecalho = {"Authorization": "Bearer " + _token(os.environ["TENANT_ID"], os.environ["CLIENT_ID"], os.environ["CLIENT_SECRET"])}
    base = f"{GRAPH}/shares/{_id_compartilhamento(link)}/driveItem"

    meta = requests.get(base, headers=cabecalho, params={"$select": "name,size,lastModifiedDateTime,lastModifiedBy"}, timeout=60)
    if meta.status_code != 200:
        raise RuntimeError(f"não foi possível ler o arquivo no SharePoint (HTTP {meta.status_code})")
    meta = meta.json()
    if meta.get("name") != cfg_pro["arquivo_base"]:
        raise RuntimeError(f"o link aponta para '{meta.get('name')}', esperado '{cfg_pro['arquivo_base']}'")

    remoto = datetime.fromisoformat(meta["lastModifiedDateTime"].replace("Z", "+00:00"))
    autor = ((meta.get("lastModifiedBy") or {}).get("user") or {}).get("displayName", "?")
    print(f"   SharePoint: alterado em {remoto.astimezone().strftime('%d/%m/%Y %H:%M:%S')} por {autor} ({meta.get('size', 0):,} bytes)")

    if destino.is_file():
        local = datetime.fromtimestamp(destino.stat().st_mtime, tz=timezone.utc)
        print(f"   Local:      alterado em {local.astimezone().strftime('%d/%m/%Y %H:%M:%S')}")
        if (remoto - local).total_seconds() <= TOLERANCIA_S:
            print("   ✅ A cópia local já está em dia; nada a baixar.")
            return False
    else:
        print("   Local:      arquivo não existe")

    pasta.mkdir(parents=True, exist_ok=True)
    tmp = Path(tempfile.NamedTemporaryFile(suffix=".xlsx", dir=pasta, delete=False).name)
    try:
        with requests.get(f"{base}/content", headers=cabecalho, stream=True, timeout=120) as r:
            if r.status_code != 200:
                raise RuntimeError(f"download falhou (HTTP {r.status_code})")
            with open(tmp, "wb") as f:
                for pedaco in r.iter_content(chunk_size=1 << 20):
                    f.write(pedaco)

        # Não substitui a cópia boa por arquivo quebrado ou sem as abas que a carga exige.
        if not zipfile.is_zipfile(tmp):
            raise RuntimeError("o arquivo baixado não é um xlsx válido")
        from openpyxl import load_workbook
        wb = load_workbook(tmp, read_only=True)
        abas = set(wb.sheetnames)
        wb.close()
        exigidas = {cfg_pro["aba_base_ativos"], cfg_pro["aba_base_rotatividade"]}
        if not exigidas <= abas:
            raise RuntimeError(f"o arquivo baixado não tem as abas {sorted(exigidas - abas)}")

        if destino.is_file():
            backups = pasta / "BACKUPS"
            backups.mkdir(exist_ok=True)
            copia = backups / f"{datetime.now().strftime('%Y_%m_%d_%H%M%S')}_{destino.name}"
            shutil.copy2(destino, copia)
            print(f"   📦 Versão anterior guardada em {copia}")
        os.replace(tmp, destino)
        os.utime(destino, (remoto.timestamp(), remoto.timestamp()))  # data local = data da fonte (comparação estável)
    finally:
        if tmp.exists():
            tmp.unlink()

    print(f"   ✅ Base atualizada a partir do SharePoint: {destino}")
    return True


if __name__ == "__main__":
    try:
        sincronizar_base_procampo()
    except Exception as erro:
        print(f"❌ {erro}")
        sys.exit(1)
